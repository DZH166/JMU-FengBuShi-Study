from pathlib import Path
import json, hashlib
from pptx import Presentation
ROOT = Path(__file__).resolve().parents[1]
out = ROOT / 'evidence' / 'sources'
out.mkdir(parents=True, exist_ok=True)
inventory=[]
for file in ROOT.parent.glob('*.pptx'):
    prs=Presentation(file)
    slides=[]
    for n, slide in enumerate(prs.slides, 1):
        chunks=[]
        def walk(shapes):
            for s in shapes:
                if s.has_text_frame:
                    chunks.append(s.text)
                if s.has_table:
                    chunks.append('\n'.join(' | '.join(c.text for c in row.cells) for row in s.table.rows))
                if s.shape_type == 6:
                    walk(s.shapes)
        walk(slide.shapes)
        slides.append({'slide':n,'text':'\n'.join(chunks)})
    data={'file':file.name,'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'slideCount':len(slides),'slides':slides}
    (out / (file.stem+'.json')).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    (out / (file.stem+'.txt')).write_text('\n\n'.join(f'--- Slide {s["slide"]} ---\n{s["text"]}' for s in slides),encoding='utf-8')
    inventory.append({'file':file.name,'slides':len(slides),'characters':sum(len(s['text']) for s in slides)})
(out/'inventory.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(inventory,ensure_ascii=False,indent=2))
